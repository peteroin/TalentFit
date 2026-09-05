from pathlib import Path

import pymupdf
from docx import Document


SUPPORTED_EXTENSIONS = {".pdf", ".docx"}


def extract_text_from_pdf(file_path: str) -> str:
    """Extract text from a PDF file."""

    text_parts = []

    document = pymupdf.open(file_path)

    try:
        for page in document:
            page_text = page.get_text("text")

            if page_text:
                text_parts.append(page_text)
    finally:
        document.close()

    return "\n".join(text_parts)


def extract_text_from_docx(file_path: str) -> str:
    """Extract text from a DOCX file."""

    document = Document(file_path)

    text_parts = []

    # Extract normal paragraphs
    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            text_parts.append(paragraph.text)

    # Extract tables
    for table in document.tables:
        for row in table.rows:
            row_text = []

            for cell in row.cells:
                cell_text = cell.text.strip()

                if cell_text:
                    row_text.append(cell_text)

            if row_text:
                text_parts.append(" | ".join(row_text))

    return "\n".join(text_parts)


def clean_text(text: str) -> str:
    """Clean extracted text."""

    lines = []

    for line in text.splitlines():
        line = line.strip()

        if line:
            lines.append(line)

    return "\n".join(lines)


def extract_resume_text(file_path: str) -> str:
    """Extract and clean text from a PDF or DOCX resume."""

    path = Path(file_path)

    extension = path.suffix.lower()

    if extension not in SUPPORTED_EXTENSIONS:
        raise ValueError(
            f"Unsupported file type: {extension}. "
            "Supported types: PDF and DOCX."
        )

    if extension == ".pdf":
        text = extract_text_from_pdf(file_path)

    else:
        text = extract_text_from_docx(file_path)

    return clean_text(text)